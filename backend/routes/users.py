import io
import os
from flask import Blueprint, request, jsonify, send_file
from flask_jwt_extended import jwt_required, get_jwt_identity
from werkzeug.security import generate_password_hash
from openpyxl import Workbook, load_workbook
from models import db, User

users_bp = Blueprint("users", __name__)


def require_admin():
    user_id = int(get_jwt_identity())
    user = User.query.get(user_id)
    if not user or not user.is_admin:
        return None
    return user


@users_bp.route("/api/users", methods=["GET"])
@jwt_required()
def list_users():
    admin = require_admin()
    if not admin:
        return jsonify({"error": "无权限"}), 403

    keyword = request.args.get("keyword", "").strip()
    page = request.args.get("page", 1, type=int)
    page_size = request.args.get("page_size", 20, type=int)

    query = User.query
    if keyword:
        query = query.filter(
            db.or_(
                User.username.contains(keyword),
                User.name.contains(keyword),
                User.department.contains(keyword),
            )
        )

    total = query.count()
    users = query.order_by(User.id.desc()).offset((page - 1) * page_size).limit(page_size).all()

    return jsonify({
        "items": [u.to_dict() for u in users],
        "total": total,
        "page": page,
        "page_size": page_size,
    })


@users_bp.route("/api/users", methods=["POST"])
@jwt_required()
def create_user():
    admin = require_admin()
    if not admin:
        return jsonify({"error": "无权限"}), 403

    data = request.get_json()
    username = data.get("username", "").strip()
    name = data.get("name", "").strip()
    department = data.get("department", "").strip()
    phone = data.get("phone", "").strip()
    is_admin = data.get("is_admin", False)
    password = data.get("password", "123456")

    if not username or not name:
        return jsonify({"error": "登录账号和用户名不能为空"}), 400

    if User.query.filter_by(username=username).first():
        return jsonify({"error": "登录账号已存在"}), 400

    user = User(
        username=username,
        name=name,
        department=department,
        phone=phone,
        is_admin=is_admin,
        password_hash=generate_password_hash(password),
    )
    db.session.add(user)
    db.session.commit()
    return jsonify(user.to_dict()), 201


@users_bp.route("/api/users/<int:user_id>", methods=["PUT"])
@jwt_required()
def update_user(user_id):
    admin = require_admin()
    if not admin:
        return jsonify({"error": "无权限"}), 403

    user = User.query.get(user_id)
    if not user:
        return jsonify({"error": "用户不存在"}), 404

    data = request.get_json()
    user.name = data.get("name", user.name)
    user.department = data.get("department", user.department)
    user.phone = data.get("phone", user.phone)
    user.is_admin = data.get("is_admin", user.is_admin)
    user.is_active = data.get("is_active", user.is_active)

    if "password" in data and data["password"]:
        user.password_hash = generate_password_hash(data["password"])

    db.session.commit()
    return jsonify(user.to_dict())


@users_bp.route("/api/users/<int:user_id>", methods=["DELETE"])
@jwt_required()
def delete_user(user_id):
    admin = require_admin()
    if not admin:
        return jsonify({"error": "无权限"}), 403

    user = User.query.get(user_id)
    if not user:
        return jsonify({"error": "用户不存在"}), 404

    current_user_id = int(get_jwt_identity())
    if user.id == current_user_id:
        return jsonify({"error": "不能删除自己"}), 400

    db.session.delete(user)
    db.session.commit()
    return jsonify({"message": "删除成功"})


@users_bp.route("/api/users/import/template", methods=["GET"])
@jwt_required()
def download_import_template():
    admin = require_admin()
    if not admin:
        return jsonify({"error": "无权限"}), 403

    wb = Workbook()
    ws = wb.active
    ws.title = "用户导入模板"
    ws.append(["登录账号", "用户名", "所属部门", "密码"])

    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 15
    ws.column_dimensions["C"].width = 20
    ws.column_dimensions["D"].width = 15

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    return send_file(
        output,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name="用户导入模板.xlsx",
    )


@users_bp.route("/api/users/import", methods=["POST"])
@jwt_required()
def import_users():
    admin = require_admin()
    if not admin:
        return jsonify({"error": "无权限"}), 403

    if "file" not in request.files:
        return jsonify({"error": "请选择文件"}), 400

    file = request.files["file"]
    if not file.filename.endswith((".xlsx", ".xls")):
        return jsonify({"error": "请上传Excel文件(.xlsx)"}), 400

    try:
        wb = load_workbook(file, read_only=True)
        ws = wb.active

        rows = list(ws.iter_rows(min_row=2, values_only=True))
        if not rows:
            return jsonify({"error": "文件为空"}), 400

        created = 0
        skipped = 0
        errors = []

        for i, row in enumerate(rows, start=2):
            username = str(row[0]).strip() if row[0] else ""
            name = str(row[1]).strip() if row[1] else ""
            department = str(row[2]).strip() if row[2] else ""
            password = str(row[3]).strip() if row[3] else "123456"

            if not username or not name:
                errors.append(f"第{i}行: 登录账号和用户名不能为空")
                skipped += 1
                continue

            if User.query.filter_by(username=username).first():
                skipped += 1
                continue

            user = User(
                username=username,
                name=name,
                department=department,
                password_hash=generate_password_hash(password),
            )
            db.session.add(user)
            created += 1

        db.session.commit()
        wb.close()

        return jsonify({
            "message": f"导入完成: 成功{created}条, 跳过{skipped}条",
            "created": created,
            "skipped": skipped,
            "errors": errors,
        })
    except Exception as e:
        return jsonify({"error": f"文件解析失败: {str(e)}"}), 400
